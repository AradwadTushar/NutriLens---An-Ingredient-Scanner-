import io
import time
from datetime import datetime

# ReportLab imports
from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

# python-docx imports
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def generate_pdf_report(data):
    """
    Generates a publication-quality clinical food intelligence PDF report.
    Returns a BytesIO buffer.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=40,
        rightMargin=40,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#2E1065')
    )

    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#6B7280')
    )

    section_heading = ParagraphStyle(
        'SectionHeading',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#1E1B4B'),
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#374151')
    )

    body_bold = ParagraphStyle(
        'BodyBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#111827')
    )

    story = []

    # 1. Header Banner
    date_str = datetime.now().strftime("%B %d, %Y - %I:%M %p")
    report_id = f"NL-{int(time.time())}"

    header_data = [
        [
            Paragraph("<b>NutriLens™</b><br/><font size=8 color='#6366F1'>Food Intelligence & Dietary Safety Dossier</font>", title_style),
            Paragraph(f"<b>Report ID:</b> #{report_id}<br/><b>Generated:</b> {date_str}<br/><b>Verification:</b> Automated Clinical Audit", subtitle_style)
        ]
    ]

    header_table = Table(header_data, colWidths=[300, 215])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ALIGN', (1, 0), (1, 0), 'RIGHT'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#4F46E5'), spaceAfter=14))

    # 2. Executive Verdict Banner
    verdict = data.get('verdict', 'Caution')
    verdict_colors = {
        'Healthy': {'bg': '#ECFDF5', 'border': '#059669', 'text': '#065F46'},
        'Caution': {'bg': '#FFFBEB', 'border': '#D97706', 'text': '#92400E'},
        'Not recommended': {'bg': '#FEF2F2', 'border': '#DC2626', 'text': '#991B1B'}
    }
    cfg = verdict_colors.get(verdict, verdict_colors['Caution'])

    verdict_text = f"<b>CLINICAL VERDICT: {verdict.upper()}</b>"
    verdict_table = Table(
        [[Paragraph(f"<font size=12 color='{cfg['text']}'>{verdict_text}</font>", body_bold)]],
        colWidths=[515]
    )
    verdict_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor(cfg['bg'])),
        ('BOX', (0, 0), (-1, -1), 1.5, colors.HexColor(cfg['border'])),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (0, 0), (-1, -1), 12),
        ('RIGHTPADDING', (0, 0), (-1, -1), 12),
    ]))
    story.append(verdict_table)
    story.append(Spacer(1, 14))

    # 3. Dietary Summary & Allergy Insights
    summary = data.get('summary', {})
    safe_text = summary.get('safe', 'No specific benefits identified.') or 'No specific benefits identified.'
    warning_text = summary.get('warning', 'No notable warnings.') or 'No notable warnings.'
    avoid_text = summary.get('avoid', 'None specified.') or 'None specified.'

    summary_rows = [
        [
            Paragraph("<b>Classification</b>", body_bold),
            Paragraph("<b>Clinical Observations & Dietary Impact</b>", body_bold)
        ],
        [
            Paragraph("<font color='#059669'><b>Positive Factors</b></font>", body_style),
            Paragraph(safe_text, body_style)
        ],
        [
            Paragraph("<font color='#D97706'><b>Cautions</b></font>", body_style),
            Paragraph(warning_text, body_style)
        ],
        [
            Paragraph("<font color='#DC2626'><b>Allergen / Avoid</b></font>", body_style),
            Paragraph(avoid_text, body_style)
        ]
    ]

    summary_table = Table(summary_rows, colWidths=[120, 395])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 14))

    # 4. Nutrition Facts Panel
    nutrition = data.get('nutrition', {})
    nutri_items = [
        ("Calories", nutrition.get('calories', 'N/A')),
        ("Protein", nutrition.get('protein', 'N/A')),
        ("Carbohydrates", nutrition.get('carbs', 'N/A')),
        ("Sugars", nutrition.get('sugar', 'N/A')),
        ("Total Fat", nutrition.get('fat', 'N/A')),
        ("Sodium", nutrition.get('sodium', 'N/A')),
        ("Fiber", nutrition.get('fiber', 'N/A')),
        ("Calcium", nutrition.get('calcium', 'N/A')),
        ("Magnesium", nutrition.get('magnesium', 'N/A')),
        ("Potassium", nutrition.get('potassium', 'N/A')),
    ]

    # Filter out items that are truly empty
    active_nutri = [(k, v) for k, v in nutri_items if v and str(v).strip()]
    if active_nutri:
        story.append(Paragraph("<b>Nutrition Facts Profile</b>", section_heading))

        # Format into a 2-column key-value grid (4 cells per row: Label | Val | Label | Val)
        nutri_grid = []
        for i in range(0, len(active_nutri), 2):
            k1, v1 = active_nutri[i]
            k2, v2 = active_nutri[i+1] if i+1 < len(active_nutri) else ("", "")
            nutri_grid.append([
                Paragraph(f"<b>{k1}</b>", body_style),
                Paragraph(str(v1), body_style),
                Paragraph(f"<b>{k2}</b>" if k2 else "", body_style),
                Paragraph(str(v2) if v2 else "", body_style)
            ])

        if nutri_grid:
            nutri_table = Table(nutri_grid, colWidths=[120, 137, 120, 138])
            nutri_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
                ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#E2E8F0')),
                ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
                ('TOPPADDING', (0, 0), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                ('LEFTPADDING', (0, 0), (-1, -1), 8),
                ('RIGHTPADDING', (0, 0), (-1, -1), 8),
            ]))
            story.append(nutri_table)
            story.append(Spacer(1, 14))

    # 5. Ingredient Safety Matrix
    ingredients = data.get('ingredients', [])
    if ingredients:
        story.append(Paragraph("<b>Analyzed Ingredients Safety Matrix</b>", section_heading))

        ing_rows = [
            [
                Paragraph("<b>Ingredient Name</b>", body_bold),
                Paragraph("<b>Status</b>", body_bold),
                Paragraph("<b>Safety Insight / Clinical Role</b>", body_bold)
            ]
        ]

        for ing in ingredients:
            name = ing.get('name', 'Unknown')
            status = ing.get('status', 'Moderate')
            desc = ing.get('description', 'Standard food ingredient.') or 'Standard food ingredient.'

            status_color = '#059669' if status == 'Good' else '#DC2626' if status == 'Risky' else '#D97706'
            ing_rows.append([
                Paragraph(f"<b>{name}</b>", body_style),
                Paragraph(f"<font color='{status_color}'><b>{status}</b></font>", body_style),
                Paragraph(desc, body_style)
            ])

        ing_table = Table(ing_rows, colWidths=[150, 75, 290])
        ing_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))
        story.append(ing_table)

    # 6. Regulatory Standards Disclaimer Footer
    story.append(Spacer(1, 16))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#94A3B8'), spaceAfter=8))
    footer_text = (
        "<b>NutriLens Dietary Audit Standards:</b> This document is compiled through automated deep-learning optical character "
        "recognition and biochemical nutritional reference cross-referencing against FDA, WHO, and FSSAI standards. "
        "Consult certified nutrition specialists for medical or metabolic condition management."
    )
    story.append(Paragraph(footer_text, subtitle_style))

    doc.build(story)
    buffer.seek(0)
    return buffer


def generate_docx_report(data):
    """
    Generates a professionally structured Microsoft Word (.docx) document.
    Returns a BytesIO buffer.
    """
    doc = Document()
    buffer = io.BytesIO()

    # Document Title
    title = doc.add_heading("NutriLens Food Intelligence Dossier", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.add_run(f"Report ID: #NL-{int(time.time())}  |  Generated: {datetime.now().strftime('%B %d, %Y - %I:%M %p')}\n").italic = True
    p_meta.add_run("Automated Clinical Dietary & Ingredient Safety Audit").font.color.rgb = RGBColor(99, 102, 241)

    # Verdict Box
    verdict = data.get('verdict', 'Caution')
    p_verdict = doc.add_paragraph()
    run_v = p_verdict.add_run(f"EXECUTIVE HEALTH VERDICT: {verdict.upper()}\n")
    run_v.bold = True
    run_v.font.size = Pt(14)
    if verdict == 'Healthy':
        run_v.font.color.rgb = RGBColor(5, 150, 105)
    elif verdict == 'Risky' or verdict == 'Not recommended':
        run_v.font.color.rgb = RGBColor(220, 38, 38)
    else:
        run_v.font.color.rgb = RGBColor(217, 119, 6)

    # Summary Section
    summary = data.get('summary', {})
    doc.add_heading("1. Dietary Assessment & Allergy Alerts", level=1)
    p_safe = doc.add_paragraph(style='List Bullet')
    p_safe.add_run("Positive Factors: ").bold = True
    p_safe.add_run(summary.get('safe', 'None specified.'))

    p_warn = doc.add_paragraph(style='List Bullet')
    p_warn.add_run("Cautions & Disclaimers: ").bold = True
    p_warn.add_run(summary.get('warning', 'None specified.'))

    p_avoid = doc.add_paragraph(style='List Bullet')
    p_avoid.add_run("Critical Warnings / Allergen Avoidance: ").bold = True
    p_avoid.add_run(summary.get('avoid', 'None specified.'))

    # Nutrition Facts Section
    nutrition = data.get('nutrition', {})
    if nutrition:
        doc.add_heading("2. Nutrition Facts Profile", level=1)
        table_nutri = doc.add_table(rows=1, cols=2)
        table_nutri.style = 'Light Shading Accent 1'
        hdr_cells = table_nutri.rows[0].cells
        hdr_cells[0].text = "Nutrient"
        hdr_cells[1].text = "Amount Per Serving"

        nutri_labels = {
            "calories": "Calories",
            "protein": "Protein",
            "carbs": "Carbohydrates",
            "sugar": "Sugars",
            "fat": "Total Fat",
            "sodium": "Sodium",
            "fiber": "Dietary Fiber",
            "calcium": "Calcium",
            "magnesium": "Magnesium",
            "potassium": "Potassium"
        }

        for key, label in nutri_labels.items():
            val = nutrition.get(key)
            if val and str(val).strip():
                row_cells = table_nutri.add_row().cells
                row_cells[0].text = label
                row_cells[1].text = str(val)

    # Ingredients Matrix Section
    ingredients = data.get('ingredients', [])
    if ingredients:
        doc.add_heading("3. Ingredient Safety & Chemical Matrix", level=1)
        table_ing = doc.add_table(rows=1, cols=3)
        table_ing.style = 'Medium Shading 1 Accent 1'
        hdr = table_ing.rows[0].cells
        hdr[0].text = "Ingredient"
        hdr[1].text = "Safety Rating"
        hdr[2].text = "Clinical Role & Impact"

        for ing in ingredients:
            row = table_ing.add_row().cells
            row[0].text = ing.get('name', 'Unknown')
            row[1].text = ing.get('status', 'Moderate')
            row[2].text = ing.get('description', 'Standard food ingredient.') or 'Standard food ingredient.'

    # Regulatory standards
    doc.add_heading("4. Regulatory Standards Note", level=1)
    p_footer = doc.add_paragraph()
    p_footer.add_run(
        "NutriLens dietary evaluations are cross-referenced against FDA, WHO, and FSSAI published nutritional reference guidelines. "
        "Generated automatically for dietary safety intelligence."
    ).italic = True

    doc.save(buffer)
    buffer.seek(0)
    return buffer

